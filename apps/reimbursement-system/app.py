# -*- coding: utf-8 -*-
"""发票与付款系统 - 后端 (Flask + SQLite, 多角色登录 + 审核流)

角色：
  admin     管理员  —— 用户管理、全部发票、批量状态、入账、导出
  approver  审核人  —— 查看待审核发票，通过/驳回（附审核意见），导出
  claimant  报销人  —— 录入/上传本人发票（草稿→提交），只看本人记录，可撤回

发票状态流：草稿 → 待审核 → 已通过 → 已入账
                     └→ 已驳回 →(修改后重新提交)→ 待审核

公司抬头（companies）：管理员维护本公司若干开票抬头；发票的购买方/销售方与之
匹配后自动判定 进项（我方为购买方）/ 销项（我方为销售方），用于进销项增值税统计。
付款：发票可登记付款日期与累计已付金额，实时计算未付差额与付款状态。
"""
import os
import re
import secrets
import shutil
import sqlite3
import time
from datetime import datetime, timedelta
from functools import wraps

from flask import (Flask, request, jsonify, session, send_from_directory,
                   send_file, after_this_request, make_response)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

from invoice_parser import parse_invoice_pdf, ParseError
from classifier import classify, all_subjects

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
DB_PATH = os.path.join(BASE_DIR, "reimbursement.db")
os.makedirs(UPLOAD_DIR, exist_ok=True)

ROLES = {"admin": "管理员", "approver": "审核人", "claimant": "报销人"}

app = Flask(__name__, static_folder="static", static_url_path="")
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024
app.permanent_session_lifetime = timedelta(days=7)

# 持久化 secret key（重启不掉线；文件缺失或为空时自动重建）
_key_file = os.path.join(BASE_DIR, ".secret_key")
app.secret_key = None
if os.path.exists(_key_file):
    try:
        with open(_key_file, "rb") as f:
            app.secret_key = f.read().decode("utf-8", "ignore").strip() or None
    except OSError:
        app.secret_key = None
if not app.secret_key:
    app.secret_key = secrets.token_hex(32)
    try:
        with open(_key_file, "wb") as f:
            f.write(app.secret_key.encode("utf-8"))
    except OSError:
        pass  # 写不进就仅内存使用（重启后登录态失效，不影响功能）


# ---------------- 数据库 ----------------
def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    return conn


def now_str():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def to_num(v):
    if v in (None, "", "-"):
        return None
    try:
        return round(float(str(v).replace(",", "")), 2)
    except (ValueError, TypeError):
        return None


def _trash(path):
    """将文件移动到 .trash 目录（重命名方式，规避直接删除），需定期人工清理"""
    if not path or not os.path.exists(path):
        return
    trash_dir = os.path.join(UPLOAD_DIR, ".trash")
    os.makedirs(trash_dir, exist_ok=True)
    try:
        os.replace(path, os.path.join(trash_dir, f"{int(time.time()*1000)}_{os.path.basename(path)}"))
    except OSError:
        pass


EDITABLE_FIELDS = [
    "invoice_no", "invoice_code", "invoice_type", "invoice_date", "buyer_name",
    "seller_name", "item_name", "amount", "tax_rate", "tax_amount", "total_amount",
    "subject_code", "subject_name", "claimant", "department", "project",
    "remark", "check_code", "settle", "payee",
]
NUMERIC_FIELDS = {"amount", "tax_amount", "total_amount"}
# 结算方式：仅做账的发票不付款、不计入欠款；收款对象：销售方 / 报销人
SETTLES = ("正常付款", "仅做账")
PAYEES = ("销售方", "报销人")


def init_db():
    conn = db()
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            invoice_no TEXT NOT NULL UNIQUE,
            invoice_code TEXT,
            invoice_type TEXT,
            invoice_date TEXT,
            buyer_name TEXT,
            seller_name TEXT,
            item_name TEXT,
            amount REAL,
            tax_rate TEXT,
            tax_amount REAL,
            total_amount REAL,
            subject_code TEXT,
            subject_name TEXT,
            claimant TEXT,
            department TEXT,
            project TEXT,
            status TEXT DEFAULT '草稿',
            remark TEXT,
            file_path TEXT,
            check_code TEXT,
            user_id INTEGER,
            created_at TEXT,
            updated_at TEXT
        );
        CREATE INDEX IF NOT EXISTS idx_invoices_date ON invoices(invoice_date);
        CREATE INDEX IF NOT EXISTS idx_invoices_subject ON invoices(subject_code);
        CREATE INDEX IF NOT EXISTS idx_invoices_status ON invoices(status);

        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL,
            display_name TEXT NOT NULL,
            role TEXT NOT NULL,
            department TEXT,
            active INTEGER DEFAULT 1,
            created_at TEXT
        );

        CREATE TABLE IF NOT EXISTS review_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            invoice_id INTEGER NOT NULL,
            reviewer_id INTEGER,
            reviewer_name TEXT,
            action TEXT NOT NULL,
            comment TEXT,
            created_at TEXT
        );
        CREATE INDEX IF NOT EXISTS idx_review_log_invoice ON review_log(invoice_id);

        CREATE TABLE IF NOT EXISTS companies (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            short_name TEXT,
            tax_no TEXT,
            active INTEGER DEFAULT 1,
            created_at TEXT
        );
        """
    )
    # 老库迁移：invoices 增加 user_id 列（必须在建 user_id 索引之前）
    cols = [r[1] for r in conn.execute("PRAGMA table_info(invoices)")]
    if "user_id" not in cols:
        conn.execute("ALTER TABLE invoices ADD COLUMN user_id INTEGER")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_invoices_user ON invoices(user_id)")
    # 老库迁移：付款与进销项方向相关列
    for col, decl in [("payment_date", "TEXT"), ("paid_amount", "REAL"),
                      ("payment_method", "TEXT"), ("payment_note", "TEXT"),
                      ("company_id", "INTEGER"), ("direction", "TEXT"),
                      ("settle", "TEXT DEFAULT '正常付款'"), ("payee", "TEXT DEFAULT '销售方'"),
                      ("cloud_path", "TEXT")]:
        if col not in cols:
            conn.execute(f"ALTER TABLE invoices ADD COLUMN {col} {decl}")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_invoices_dir ON invoices(direction)")
    # 种子账号（仅当用户表为空时创建）
    if not conn.execute("SELECT 1 FROM users LIMIT 1").fetchone():
        seed = [
            ("admin", "admin123", "系统管理员", "admin", "财务部"),
            ("shenhe", "shenhe123", "审核专员", "approver", "财务部"),
            ("baoxiao", "baoxiao123", "报销演示", "claimant", "业务部"),
        ]
        for username, pwd, name, role, dept in seed:
            conn.execute(
                "INSERT INTO users (username,password_hash,display_name,role,department,active,created_at)"
                " VALUES (?,?,?,?,?,1,?)",
                (username, generate_password_hash(pwd), name, role, dept, now_str()))
    conn.commit()
    conn.close()


init_db()


# ---------------- 公司抬头匹配（进项/销项方向） ----------------
def _norm_name(s):
    return re.sub(r"\s+", "", s or "")


def _match_company(comps, buyer, seller):
    """按公司抬头匹配发票购销方，返回 (company_id, direction)。优先匹配购买方（进项）。"""
    b, s = _norm_name(buyer), _norm_name(seller)
    if not b and not s:
        return None, None
    for c in comps:
        n = _norm_name(c["name"])
        if n and b and (n in b or b in n):
            return c["id"], "进项"
    for c in comps:
        n = _norm_name(c["name"])
        if n and s and (n in s or s in n):
            return c["id"], "销项"
    return None, None


def compute_direction(conn, buyer, seller):
    comps = conn.execute("SELECT id, name FROM companies WHERE active=1").fetchall()
    return _match_company(comps, buyer, seller)


def recompute_all_directions(conn):
    """公司抬头变化后，重算全部发票的进项/销项归属"""
    comps = conn.execute("SELECT id, name FROM companies WHERE active=1").fetchall()
    for r in conn.execute("SELECT id, buyer_name, seller_name FROM invoices").fetchall():
        cid, d = _match_company(comps, r["buyer_name"], r["seller_name"])
        conn.execute("UPDATE invoices SET company_id=?, direction=? WHERE id=?", (cid, d, r["id"]))


def payment_info(total, paid):
    """根据价税合计与累计已付金额，返回 (付款状态, 未付差额)"""
    if total is None:
        return "-", None
    paid = paid or 0
    bal = round(total - paid, 2)
    if paid <= 0.005:
        return "未付款", bal
    if bal > 0.005:
        return "部分付款", bal
    return "已付清", 0.0


# 启动时确保全部发票的进/销项方向与现有公司抬头一致
_c = db()
recompute_all_directions(_c)
_c.commit()
_c.close()


# ---------------- 认证 ----------------
def current_user():
    uid = session.get("uid")
    if not uid:
        return None
    conn = db()
    u = conn.execute("SELECT * FROM users WHERE id=? AND active=1", (uid,)).fetchone()
    conn.close()
    return dict(u) if u else None


def login_required(*roles):
    def deco(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            u = current_user()
            if not u:
                return jsonify({"ok": False, "msg": "请先登录"}), 401
            if roles and u["role"] not in roles:
                return jsonify({"ok": False, "msg": "没有权限执行此操作"}), 403
            return fn(*args, **kwargs)
        return wrapper
    return deco


def add_log(conn, invoice_id, user, action, comment=None):
    conn.execute(
        "INSERT INTO review_log (invoice_id,reviewer_id,reviewer_name,action,comment,created_at)"
        " VALUES (?,?,?,?,?,?)",
        (invoice_id, user["id"], user["display_name"], action, comment, now_str()))


def get_invoice(conn, iid):
    return conn.execute("SELECT * FROM invoices WHERE id=?", (iid,)).fetchone()


# ---------------- 登录用户名列表（公开，仅启用账号的显示名） ----------------
@app.get("/api/login-users")
def login_users():
    conn = db()
    rows = conn.execute(
        "SELECT username, display_name FROM users WHERE active=1 ORDER BY id").fetchall()
    conn.close()
    return jsonify({"ok": True, "rows": [dict(r) for r in rows]})


# ---------------- 登录 / 登出 / 当前用户 ----------------
@app.post("/api/login")
def login():
    body = request.get_json(silent=True) or {}
    username = (body.get("username") or "").strip()
    password = body.get("password") or ""
    if not username or not password:
        return jsonify({"ok": False, "msg": "请输入用户名和密码"}), 400
    conn = db()
    u = conn.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
    conn.close()
    if not u or not check_password_hash(u["password_hash"], password):
        return jsonify({"ok": False, "msg": "用户名或密码错误"}), 401
    if not u["active"]:
        return jsonify({"ok": False, "msg": "该账号已被停用，请联系管理员"}), 403
    session["uid"] = u["id"]
    session.permanent = True
    return jsonify({"ok": True, "user": {
        "id": u["id"], "username": u["username"], "display_name": u["display_name"],
        "role": u["role"], "role_name": ROLES[u["role"]], "department": u["department"]}})


@app.post("/api/logout")
def logout():
    session.clear()
    return jsonify({"ok": True})


@app.get("/api/me")
def me():
    u = current_user()
    if not u:
        return jsonify({"ok": False, "msg": "未登录"}), 401
    return jsonify({"ok": True, "user": {
        "id": u["id"], "username": u["username"], "display_name": u["display_name"],
        "role": u["role"], "role_name": ROLES[u["role"]], "department": u["department"]}})


# ---------------- 用户管理（仅管理员） ----------------
@app.get("/api/users")
@login_required("admin")
def users_list():
    conn = db()
    rows = conn.execute(
        "SELECT id, username, display_name, role, department, active, created_at FROM users ORDER BY id").fetchall()
    conn.close()
    return jsonify({"ok": True, "rows": [dict(r) for r in rows]})


@app.post("/api/users")
@login_required("admin")
def users_create():
    b = request.get_json(silent=True) or {}
    username = (b.get("username") or "").strip()
    password = b.get("password") or ""
    display_name = (b.get("display_name") or "").strip()
    role = b.get("role")
    department = (b.get("department") or "").strip() or None
    if not username or not password or not display_name:
        return jsonify({"ok": False, "msg": "用户名/密码/姓名必填"}), 400
    if role not in ROLES:
        return jsonify({"ok": False, "msg": "角色不合法"}), 400
    if len(password) < 6:
        return jsonify({"ok": False, "msg": "密码至少 6 位"}), 400
    conn = db()
    if conn.execute("SELECT 1 FROM users WHERE username=?", (username,)).fetchone():
        conn.close()
        return jsonify({"ok": False, "msg": f"用户名 {username} 已存在"}), 409
    conn.execute(
        "INSERT INTO users (username,password_hash,display_name,role,department,active,created_at)"
        " VALUES (?,?,?,?,?,1,?)",
        (username, generate_password_hash(password), display_name, role, department, now_str()))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.put("/api/users/<int:uid>")
@login_required("admin")
def users_update(uid):
    b = request.get_json(silent=True) or {}
    conn = db()
    u = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
    if not u:
        conn.close()
        return jsonify({"ok": False, "msg": "用户不存在"}), 404
    if "role" in b:
        if b["role"] not in ROLES:
            conn.close()
            return jsonify({"ok": False, "msg": "角色不合法"}), 400
        if uid == session["uid"] and b["role"] != "admin":
            conn.close()
            return jsonify({"ok": False, "msg": "不能降级自己的管理员角色"}), 400
        conn.execute("UPDATE users SET role=? WHERE id=?", (b["role"], uid))
    if "display_name" in b and (b["display_name"] or "").strip():
        conn.execute("UPDATE users SET display_name=? WHERE id=?", (b["display_name"].strip(), uid))
    if "department" in b:
        conn.execute("UPDATE users SET department=? WHERE id=?", ((b["department"] or "").strip() or None, uid))
    if "active" in b:
        if uid == session["uid"] and not b["active"]:
            conn.close()
            return jsonify({"ok": False, "msg": "不能停用自己的账号"}), 400
        conn.execute("UPDATE users SET active=? WHERE id=?", (1 if b["active"] else 0, uid))
    if b.get("password"):
        if len(b["password"]) < 6:
            conn.close()
            return jsonify({"ok": False, "msg": "密码至少 6 位"}), 400
        conn.execute("UPDATE users SET password_hash=? WHERE id=?",
                     (generate_password_hash(b["password"]), uid))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


# ---------------- 公司抬头管理（管理员维护，全员可读） ----------------
@app.get("/api/companies")
@login_required()
def companies_list():
    conn = db()
    rows = conn.execute(
        "SELECT c.*, (SELECT COUNT(*) FROM invoices i WHERE i.company_id=c.id) n_invoices"
        " FROM companies c ORDER BY c.id").fetchall()
    conn.close()
    return jsonify({"ok": True, "rows": [dict(r) for r in rows]})


@app.post("/api/companies")
@login_required("admin")
def companies_create():
    b = request.get_json(silent=True) or {}
    name = (b.get("name") or "").strip()
    if not name:
        return jsonify({"ok": False, "msg": "公司全称必填"}), 400
    conn = db()
    if conn.execute("SELECT 1 FROM companies WHERE name=?", (name,)).fetchone():
        conn.close()
        return jsonify({"ok": False, "msg": f"公司抬头“{name}”已存在"}), 409
    conn.execute(
        "INSERT INTO companies (name,short_name,tax_no,active,created_at) VALUES (?,?,?,?,?)",
        (name, (b.get("short_name") or "").strip() or None,
         (b.get("tax_no") or "").strip() or None, 1, now_str()))
    recompute_all_directions(conn)
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.put("/api/companies/<int:cid>")
@login_required("admin")
def companies_update(cid):
    b = request.get_json(silent=True) or {}
    conn = db()
    c = conn.execute("SELECT * FROM companies WHERE id=?", (cid,)).fetchone()
    if not c:
        conn.close()
        return jsonify({"ok": False, "msg": "公司不存在"}), 404
    if "name" in b:
        name = (b["name"] or "").strip()
        if not name:
            conn.close()
            return jsonify({"ok": False, "msg": "公司全称不能为空"}), 400
        dup = conn.execute("SELECT 1 FROM companies WHERE name=? AND id<>?", (name, cid)).fetchone()
        if dup:
            conn.close()
            return jsonify({"ok": False, "msg": f"公司抬头“{name}”已存在"}), 409
        conn.execute("UPDATE companies SET name=? WHERE id=?", (name, cid))
    if "short_name" in b:
        conn.execute("UPDATE companies SET short_name=? WHERE id=?", ((b["short_name"] or "").strip() or None, cid))
    if "tax_no" in b:
        conn.execute("UPDATE companies SET tax_no=? WHERE id=?", ((b["tax_no"] or "").strip() or None, cid))
    if "active" in b:
        conn.execute("UPDATE companies SET active=? WHERE id=?", (1 if b["active"] else 0, cid))
    recompute_all_directions(conn)
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.delete("/api/companies/<int:cid>")
@login_required("admin")
def companies_delete(cid):
    conn = db()
    c = conn.execute("SELECT * FROM companies WHERE id=?", (cid,)).fetchone()
    if not c:
        conn.close()
        return jsonify({"ok": False, "msg": "公司不存在"}), 404
    conn.execute("DELETE FROM companies WHERE id=?", (cid,))
    recompute_all_directions(conn)
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "msg": f"已删除“{c['name']}”，相关发票的进销项方向已重新匹配"})


@app.delete("/api/users/<int:uid>")
@login_required("admin")
def users_delete(uid):
    """删除用户：本人不可删；名下发票保留在台账中（claimant 等字段不受影响）"""
    if uid == session.get("uid"):
        return jsonify({"ok": False, "msg": "不能删除自己的账号"}), 400
    conn = db()
    u = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
    if not u:
        conn.close()
        return jsonify({"ok": False, "msg": "用户不存在"}), 404
    n = conn.execute("SELECT COUNT(*) c FROM invoices WHERE user_id=?", (uid,)).fetchone()["c"]
    conn.execute("DELETE FROM users WHERE id=?", (uid,))
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "msg": f"已删除用户 {u['display_name']}（{u['username']}），名下 {n} 张发票保留在台账中"})


# ---------------- 批量提交（报销人/管理员，仅草稿与已驳回） ----------------
@app.post("/api/invoices/batch-submit")
@login_required("claimant", "admin")
def batch_submit():
    u = current_user()
    body = request.get_json(silent=True) or {}
    ids = body.get("ids", [])
    if not ids:
        return jsonify({"ok": False, "msg": "参数错误"}), 400
    conn = db()
    updated = skipped = 0
    for iid in ids:
        r = get_invoice(conn, iid)
        if not r or r["status"] not in ("草稿", "已驳回"):
            skipped += 1
            continue
        if u["role"] == "claimant" and r["user_id"] != u["id"]:
            skipped += 1
            continue
        conn.execute("UPDATE invoices SET status='待审核', updated_at=? WHERE id=?", (now_str(), iid))
        add_log(conn, iid, u, "提交", "批量提交")
        updated += 1
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "updated": updated, "skipped": skipped})


# ---------------- 页面 / 静态 ----------------
@app.get("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


# ---------------- 科目列表 ----------------
@app.get("/api/subjects")
@login_required()
def subjects():
    return jsonify(all_subjects())


# ---------------- 解析预览（不入库） ----------------
@app.post("/api/parse")
@login_required("claimant", "admin")
def parse_preview():
    f = request.files.get("file")
    if not f or not f.filename.lower().endswith(".pdf"):
        return jsonify({"ok": False, "msg": "请上传 PDF 文件"}), 400
    tmp = os.path.join(UPLOAD_DIR, f"~preview_{int(time.time()*1000)}_{secure_filename(f.filename)}")
    f.save(tmp)
    try:
        data = parse_invoice_pdf(tmp)
        code, name, hit = classify(data.get("item_name"), data.get("seller_name"), data.get("tax_rate"))
        data["subject_code"], data["subject_name"] = code, name
        data["suggest_reason"] = hit
        data["ok"] = True
        return jsonify(data)
    except ParseError as e:
        return jsonify({"ok": False, "msg": str(e)}), 422
    except Exception as e:
        return jsonify({"ok": False, "msg": f"解析失败: {e}"}), 500
    finally:
        _trash(tmp)


# ---------------- 上传并入库（发票号唯一） ----------------
@app.post("/api/upload")
@login_required("claimant", "admin")
def upload():
    u = current_user()
    f = request.files.get("file")
    form = request.form
    if not f or not f.filename.lower().endswith(".pdf"):
        return jsonify({"ok": False, "msg": "请上传 PDF 文件"}), 400

    stage = os.path.join(UPLOAD_DIR, f"~stage_{int(time.time()*1000)}_{secure_filename(f.filename)}")
    f.save(stage)

    try:
        invoice_no = (form.get("invoice_no") or "").strip()
        parsed = None
        if not invoice_no:
            try:
                parsed = parse_invoice_pdf(stage)
                invoice_no = parsed["invoice_no"]
            except ParseError as e:
                _trash(stage)
                return jsonify({"ok": False, "msg": f"PDF 解析失败：{e}"}), 422
        if not invoice_no:
            _trash(stage)
            return jsonify({"ok": False, "msg": "无法确定发票号码，请手工录入发票号。"}), 422

        conn = db()
        dup = conn.execute("SELECT id FROM invoices WHERE invoice_no = ?", (invoice_no,)).fetchone()
        if dup:
            conn.close()
            _trash(stage)
            return jsonify({"ok": False, "duplicate": True,
                            "msg": f"发票号 {invoice_no} 已存在（台账编号 #{dup['id']}），不可重复录入。"}), 409

        safe_no = re.sub(r"[^\w-]", "_", invoice_no)
        fpath = os.path.join(UPLOAD_DIR, f"{safe_no}.pdf")
        if os.path.exists(fpath):
            _trash(fpath)
        os.replace(stage, fpath)

        if parsed is None:
            try:
                parsed = parse_invoice_pdf(fpath)
            except Exception:
                parsed = None
        p = parsed or {}
        rec = {
            "invoice_no": invoice_no,
            "invoice_code": form.get("invoice_code") or p.get("invoice_code"),
            "invoice_type": form.get("invoice_type") or p.get("invoice_type"),
            "invoice_date": form.get("invoice_date") or p.get("invoice_date"),
            "buyer_name": form.get("buyer_name") or p.get("buyer_name"),
            "seller_name": form.get("seller_name") or p.get("seller_name"),
            "item_name": form.get("item_name") or p.get("item_name"),
            "amount": to_num(form.get("amount") or p.get("amount")),
            "tax_rate": form.get("tax_rate") or p.get("tax_rate"),
            "tax_amount": to_num(form.get("tax_amount") or p.get("tax_amount")),
            "total_amount": to_num(form.get("total_amount") or p.get("total_amount")),
            "subject_code": form.get("subject_code"),
            "subject_name": form.get("subject_name"),
            "claimant": u["display_name"] if u["role"] == "claimant" else (form.get("claimant") or u["display_name"]),
            "department": u["department"] if (u["role"] == "claimant" and u["department"]) else form.get("department"),
            "project": form.get("project") or None,
            "remark": form.get("remark") or None,
            "check_code": p.get("check_code"),
            "file_path": os.path.basename(fpath),
            "user_id": u["id"],
            # 提交=1 直接进待审核；否则存草稿
            "status": "待审核" if form.get("submit") == "1" else "草稿",
        }
        if not rec["subject_code"]:
            code, name, _ = classify(rec["item_name"], rec["seller_name"], rec["tax_rate"])
            rec["subject_code"], rec["subject_name"] = code, name
        rec["created_at"] = rec["updated_at"] = now_str()

        cols = ",".join(rec)
        qs = ",".join("?" for _ in rec)
        cur = conn.execute(f"INSERT INTO invoices ({cols}) VALUES ({qs})", tuple(rec.values()))
        # 按公司抬头判定进项/销项方向
        cid, direction = compute_direction(conn, rec["buyer_name"], rec["seller_name"])
        conn.execute("UPDATE invoices SET company_id=?, direction=? WHERE id=?", (cid, direction, cur.lastrowid))
        company_short = None
        if cid:
            crow = conn.execute("SELECT short_name, name FROM companies WHERE id=?", (cid,)).fetchone()
            if crow:
                company_short = crow["short_name"] or crow["name"]
        add_log(conn, cur.lastrowid, u, "提交" if rec["status"] == "待审核" else "草稿")
        conn.commit()
        conn.close()
        return jsonify({"ok": True, "id": cur.lastrowid, "invoice_no": invoice_no,
                        "status": rec["status"], "parsed": bool(parsed), "subject": rec["subject_name"],
                        "company_id": cid, "direction": direction, "company_short": company_short,
                        "file_path": rec["file_path"]})
    except Exception as e:
        _trash(stage)
        return jsonify({"ok": False, "msg": f"上传处理失败: {e}"}), 500


# ---------------- 外部系统提交（采购系统 → 待审核） ----------------
# 与采购一体化系统（purchase-integrated/invoice.html）约定：
#   POST /api/invoices/external  JSON + 请求头 X-Reim-Token
# 供应商在采购系统录入发票后点「提交」，此处入库为「待审核」并记审核日志。
# 联网部署时请更换 EXTERNAL_TOKEN，并把 Allow-Origin 收紧为采购系统的域名。
EXTERNAL_TOKEN = "purchase-link-2026"


def _external_cors(resp):
    resp.headers["Access-Control-Allow-Origin"] = "*"
    resp.headers["Access-Control-Allow-Headers"] = "Content-Type, X-Reim-Token"
    resp.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS"
    return resp


@app.route("/api/invoices/external", methods=["POST", "OPTIONS"])
def invoice_external():
    if request.method == "OPTIONS":
        return _external_cors(make_response("", 204))
    if request.headers.get("X-Reim-Token", "") != EXTERNAL_TOKEN:
        return _external_cors(make_response(jsonify({"ok": False, "msg": "令牌无效"}), 401))
    d = request.get_json(silent=True) or {}
    invoice_no = str(d.get("invoice_no") or "").strip()
    if not invoice_no:
        return _external_cors(make_response(jsonify({"ok": False, "msg": "缺少发票号码"}), 422))
    try:
        conn = db()
        existed = conn.execute(
            "SELECT id, status FROM invoices WHERE invoice_no=?", (invoice_no,)).fetchone()
        rec = {
            "invoice_no": invoice_no,
            "invoice_code": d.get("invoice_code"),
            "invoice_type": d.get("invoice_type"),
            "invoice_date": d.get("invoice_date"),
            "buyer_name": d.get("buyer_name"),
            "seller_name": d.get("seller_name"),
            "item_name": d.get("item_name"),
            "amount": to_num(d.get("amount")),
            "tax_rate": d.get("tax_rate"),
            "tax_amount": to_num(d.get("tax_amount")),
            "total_amount": to_num(d.get("total_amount")),
            "subject_code": d.get("subject_code"),
            "subject_name": d.get("subject_name"),
            "claimant": (d.get("claimant") or "采购系统"),
            "department": d.get("department"),
            "project": d.get("project"),
            "status": "待审核",
            "remark": d.get("remark"),
            "check_code": d.get("check_code"),
            "settle": "正常付款",
            "payee": "销售方",
        }
        if not rec.get("subject_code"):
            code, name, _ = classify(rec["item_name"], rec["seller_name"], rec["tax_rate"])
            rec["subject_code"], rec["subject_name"] = code, name
        if existed:
            # 仅允许「已驳回」的发票由采购侧修改后重新提交；其余状态一律拒绝重复提交
            if existed["status"] != "已驳回":
                conn.close()
                return _external_cors(make_response(jsonify(
                    {"ok": False, "duplicate": True,
                     "msg": f"发票 {invoice_no} 已在报销台账中（{existed['status']}），不可重复提交"}), 409))
            iid = existed["id"]
            rec["updated_at"] = now_str()
            sets = ",".join(f"{k}=?" for k in rec)
            conn.execute(f"UPDATE invoices SET {sets} WHERE id=?", tuple(rec.values()) + (iid,))
            cid, direction = compute_direction(conn, rec["buyer_name"], rec["seller_name"])
            conn.execute("UPDATE invoices SET company_id=?, direction=? WHERE id=?",
                         (cid, direction, iid))
            add_log(conn, iid, {"id": 0, "display_name": "采购系统"}, "采购系统重新提交")
            conn.commit()
            conn.close()
            return _external_cors(make_response(jsonify(
                {"ok": True, "id": iid, "invoice_no": invoice_no,
                 "status": "待审核", "subject": rec["subject_name"], "resubmitted": True}), 200))
        rec["file_path"] = None
        rec["user_id"] = None
        rec["created_at"] = rec["updated_at"] = now_str()
        cols = ",".join(rec)
        qs = ",".join("?" for _ in rec)
        cur = conn.execute(f"INSERT INTO invoices ({cols}) VALUES ({qs})", tuple(rec.values()))
        cid, direction = compute_direction(conn, rec["buyer_name"], rec["seller_name"])
        conn.execute("UPDATE invoices SET company_id=?, direction=? WHERE id=?",
                     (cid, direction, cur.lastrowid))
        add_log(conn, cur.lastrowid, {"id": 0, "display_name": "采购系统"}, "采购系统提交")
        conn.commit()
        conn.close()
        return _external_cors(make_response(jsonify(
            {"ok": True, "id": cur.lastrowid, "invoice_no": invoice_no,
             "status": rec["status"], "subject": rec["subject_name"]}), 201))
    except Exception as e:
        return _external_cors(make_response(jsonify({"ok": False, "msg": f"入库失败: {e}"}), 500))


# ---------------- 外部系统读取（采购系统 ← 付款数据，只读） ----------------
# 与采购一体化系统（purchase-integrated/payment.html）约定：
#   GET /api/payments/external  + 请求头 X-Reim-Token（同 EXTERNAL_TOKEN）
# 把「收票付款统计（付给销售方）」与「付款记录」映射给采购侧付款管理页，
# 供供应商查看 / 查询 / 核对。只读接口，不接收任何写入。
@app.route("/api/payments/external", methods=["GET", "OPTIONS"])
def payments_external():
    if request.method == "OPTIONS":
        return _external_cors(make_response("", 204))
    if request.headers.get("X-Reim-Token", "") != EXTERNAL_TOKEN:
        return _external_cors(make_response(jsonify({"ok": False, "msg": "令牌无效"}), 401))
    seller = (request.args.get("seller") or "").strip()
    conn = db()
    # 口径与 /api/stats「收票付款统计」一致：正常付款、非驳回的进项票；
    # 欠款只统计付给销售方（供应商）的票，仅做账 / 已驳回 / 付报销人不计入
    pay_base = ("direction='进项' AND total_amount IS NOT NULL"
                " AND settle='正常付款' AND status!='已驳回'")
    seller_cond = " AND seller_name = ?" if seller else ""
    sparams = [seller] if seller else []
    groups = conn.execute(
        f"""SELECT seller_name, COUNT(*) n, SUM(total_amount) t,
                   SUM(COALESCE(paid_amount,0)) p, MAX(payment_date) last_pay
            FROM invoices WHERE {pay_base} AND payee='销售方'{seller_cond}
            GROUP BY seller_name ORDER BY t DESC""", sparams).fetchall()
    # 全部有效进项票明细（含未付款），供应商按张核对付款进度
    invs = conn.execute(
        f"""SELECT id, invoice_no, invoice_type, invoice_date, seller_name, item_name,
                   project, payee, settle, total_amount, paid_amount,
                   payment_date, payment_method, status
            FROM invoices WHERE {pay_base} AND payee='销售方'{seller_cond}
            ORDER BY seller_name, invoice_date DESC, id DESC""", sparams).fetchall()
    # 付款记录：同上进项 / 正常付款 / 付给销售方 / 非驳回口径，仅取已有付款的票，
    # 保证采购侧付款流水、欠款与供应商分组口径一致（不混入报销人/客户款项）
    recs = conn.execute(
        f"""SELECT id, invoice_no, invoice_type, invoice_date, seller_name, claimant,
                   payee, settle, project, total_amount, paid_amount,
                   payment_date, payment_method, payment_note, status
            FROM invoices WHERE {pay_base} AND payee='销售方'
                 AND paid_amount IS NOT NULL AND paid_amount > 0{seller_cond}
            ORDER BY payment_date DESC, id DESC LIMIT 2000""", sparams).fetchall()
    # 采购系统提交单的最新审核状态（以「采购系统提交/重新提交」日志为来源标记），
    # 供采购侧发票台账回写「待审核/已通过/已入账/已驳回」及驳回原因
    subs = conn.execute(
        """SELECT i.id, i.invoice_no, i.status, i.updated_at,
                  (SELECT g.comment FROM review_log g
                    WHERE g.invoice_id=i.id AND g.comment IS NOT NULL AND g.comment<>''
                    ORDER BY g.id DESC LIMIT 1) AS last_comment
             FROM invoices i
             WHERE i.id IN (SELECT DISTINCT invoice_id FROM review_log
                             WHERE action IN ('采购系统提交','采购系统重新提交'))
             ORDER BY i.updated_at DESC""").fetchall()
    conn.close()
    inv_rows = []
    for r in invs:
        d = dict(r)
        d["payment_status"], d["balance"] = payment_info(d["total_amount"], d["paid_amount"])
        inv_rows.append(d)
    rec_rows = []
    for r in recs:
        d = dict(r)
        d["payment_status"], d["balance"] = payment_info(d["total_amount"], d["paid_amount"])
        rec_rows.append(d)
    total_n = sum(g["n"] for g in groups)
    total_t = sum(g["t"] or 0 for g in groups)
    total_p = sum(g["p"] or 0 for g in groups)
    return _external_cors(make_response(jsonify({
        "ok": True,
        "generated_at": now_str(),
        "summary": {"suppliers": len(groups), "invoices": total_n,
                    "total": round(total_t, 2), "paid": round(total_p, 2),
                    "owed": round(total_t - total_p, 2)},
        "groups": [{"seller_name": g["seller_name"], "count": g["n"],
                    "total": round(g["t"] or 0, 2), "paid": round(g["p"] or 0, 2),
                    "owed": round((g["t"] or 0) - (g["p"] or 0), 2),
                    "last_pay": g["last_pay"] or ""} for g in groups],
        "invoices": inv_rows,
        "records": rec_rows,
        "submissions": [dict(r) for r in subs],
    }), 200))


# ---------------- 查询（角色范围 + 筛选 + 分页） ----------------
@app.get("/api/invoices")
@login_required()
def list_invoices():
    u = current_user()
    q = request.args
    where, params = ["1=1"], []
    if u["role"] == "claimant":
        where.append("user_id = ?")
        params.append(u["id"])
    elif u["role"] == "approver":
        where.append("status != '草稿'")
    for field in ["invoice_type", "subject_code", "status", "claimant", "department"]:
        v = (q.get(field) or "").strip()
        if v:
            where.append(f"{field} = ?")
            params.append(v)
    kw = (q.get("keyword") or "").strip()
    if kw:
        like = f"%{kw}%"
        where.append("(invoice_no LIKE ? OR seller_name LIKE ? OR item_name LIKE ? OR "
                     "buyer_name LIKE ? OR remark LIKE ? OR claimant LIKE ? OR project LIKE ?)")
        params += [like] * 7
    date_from, date_to = (q.get("date_from") or "").strip(), (q.get("date_to") or "").strip()
    if date_from:
        where.append("invoice_date >= ?")
        params.append(date_from)
    if date_to:
        where.append("invoice_date <= ?")
        params.append(date_to)

    cond = " AND ".join(where)
    sort = q.get("sort", "invoice_date")
    order = "DESC" if q.get("order", "desc") == "desc" else "ASC"
    if sort not in {"invoice_date", "total_amount", "created_at", "invoice_no", "subject_code"}:
        sort = "invoice_date"

    page = max(1, int(q.get("page", 1)))
    size = min(1000, max(1, int(q.get("page_size", 20))))
    conn = db()
    total = conn.execute(f"SELECT COUNT(*) c FROM invoices WHERE {cond}", params).fetchone()["c"]
    rows = conn.execute(
        f"""SELECT i.*,
              (SELECT g.comment FROM review_log g WHERE g.invoice_id = i.id ORDER BY g.id DESC LIMIT 1) AS last_comment,
              (SELECT g.action FROM review_log g WHERE g.invoice_id = i.id ORDER BY g.id DESC LIMIT 1) AS last_action,
              (SELECT COALESCE(NULLIF(c.short_name,''), c.name) FROM companies c WHERE c.id=i.company_id) AS company_short
            FROM invoices i WHERE {cond} ORDER BY {sort} {order}, id DESC LIMIT ? OFFSET ?""",
        params + [size, (page - 1) * size]).fetchall()
    agg = conn.execute(
        f"SELECT SUM(total_amount) t, SUM(tax_amount) x, SUM(amount) a FROM invoices WHERE {cond}",
        params).fetchone()
    conn.close()
    out_rows = []
    for r in rows:
        d = dict(r)
        d["payment_status"], d["balance"] = payment_info(d.get("total_amount"), d.get("paid_amount"))
        out_rows.append(d)
    return jsonify({"ok": True, "total": total, "page": page, "page_size": size,
                    "rows": out_rows,
                    "sum_total": agg["t"] or 0, "sum_tax": agg["x"] or 0, "sum_amount": agg["a"] or 0})


# ---------------- 编辑 ----------------
@app.put("/api/invoices/<int:iid>")
@login_required("claimant", "admin")
def update_invoice(iid):
    u = current_user()
    body = request.get_json(silent=True) or {}
    conn = db()
    r = get_invoice(conn, iid)
    if not r:
        conn.close()
        return jsonify({"ok": False, "msg": "记录不存在"}), 404
    # 已通过/已入账的发票：除管理员外任何人不可编辑
    if r["status"] in ("已通过", "已入账") and u["role"] != "admin":
        conn.close()
        return jsonify({"ok": False, "msg": "已通过审核的发票仅管理员可编辑"}), 403
    if u["role"] == "claimant":
        if r["user_id"] != u["id"]:
            conn.close()
            return jsonify({"ok": False, "msg": "只能编辑本人的发票"}), 403
        if r["status"] not in ("草稿", "已驳回"):
            conn.close()
            return jsonify({"ok": False, "msg": f"当前状态“{r['status']}”不可编辑"}), 403

    updates = {k: body[k] for k in EDITABLE_FIELDS if k in body}
    if not updates:
        conn.close()
        return jsonify({"ok": False, "msg": "没有需要更新的字段"}), 400
    if "settle" in updates and updates["settle"] not in SETTLES:
        conn.close()
        return jsonify({"ok": False, "msg": "结算方式不合法"}), 400
    if "payee" in updates and updates["payee"] not in PAYEES:
        conn.close()
        return jsonify({"ok": False, "msg": "收款对象不合法"}), 400
    for k in NUMERIC_FIELDS:
        if k in updates:
            updates[k] = to_num(updates[k])
    if "invoice_no" in updates:
        new_no = (updates["invoice_no"] or "").strip()
        if not new_no:
            conn.close()
            return jsonify({"ok": False, "msg": "发票号码不能为空"}), 400
        dup = conn.execute("SELECT id FROM invoices WHERE invoice_no=? AND id<>?", (new_no, iid)).fetchone()
        if dup:
            conn.close()
            return jsonify({"ok": False, "msg": f"发票号 {new_no} 已被台账 #{dup['id']} 占用。"}), 409
        updates["invoice_no"] = new_no
    if "subject_code" in updates and not updates.get("subject_name"):
        for c, n in all_subjects():
            if c == updates["subject_code"]:
                updates["subject_name"] = n
    updates["updated_at"] = now_str()
    sets = ",".join(f"{k}=?" for k in updates)
    conn.execute(f"UPDATE invoices SET {sets} WHERE id=?", tuple(updates.values()) + (iid,))
    # 购买方/销售方变动后重新匹配进销项方向
    if "buyer_name" in updates or "seller_name" in updates:
        r2 = get_invoice(conn, iid)
        cid, direction = compute_direction(conn, r2["buyer_name"], r2["seller_name"])
        conn.execute("UPDATE invoices SET company_id=?, direction=? WHERE id=?", (cid, direction, iid))
    # 报销人修改已驳回发票后，自动回到草稿，等待重新提交
    if u["role"] == "claimant" and r["status"] == "已驳回":
        conn.execute("UPDATE invoices SET status='草稿', updated_at=? WHERE id=?", (now_str(), iid))
    conn.commit()
    # 返回更新后的进项/销项归属，供前端把云端 PDF 迁移到对应文件夹
    nr = conn.execute(
        "SELECT i.company_id, i.direction, i.cloud_path,"
        " (SELECT COALESCE(NULLIF(c.short_name,''), c.name) FROM companies c WHERE c.id=i.company_id) company_short"
        " FROM invoices i WHERE id=?", (iid,)).fetchone()
    conn.close()
    return jsonify({"ok": True,
                    "company_id": nr["company_id"], "direction": nr["direction"],
                    "company_short": nr["company_short"], "cloud_path": nr["cloud_path"]})


# ---------------- 删除 ----------------
@app.delete("/api/invoices/<int:iid>")
@login_required("claimant", "admin")
def delete_invoice(iid):
    u = current_user()
    conn = db()
    r = get_invoice(conn, iid)
    if not r:
        conn.close()
        return jsonify({"ok": False, "msg": "记录不存在"}), 404
    # 已通过/已入账的发票：除管理员外任何人不可删除
    if r["status"] in ("已通过", "已入账") and u["role"] != "admin":
        conn.close()
        return jsonify({"ok": False, "msg": "已通过审核的发票仅管理员可删除"}), 403
    if u["role"] == "claimant":
        if r["user_id"] != u["id"] or r["status"] not in ("草稿", "已驳回"):
            conn.close()
            return jsonify({"ok": False, "msg": "只能删除本人草稿或已驳回的发票"}), 403
    fpath = r["file_path"]
    conn.execute("DELETE FROM review_log WHERE invoice_id=?", (iid,))
    conn.execute("DELETE FROM invoices WHERE id=?", (iid,))
    conn.commit()
    conn.close()
    if fpath:
        fp = os.path.join(UPLOAD_DIR, os.path.basename(fpath))
        if os.path.exists(fp):
            _trash(fp)
    return jsonify({"ok": True})


# ---------------- 云端 PDF 路径登记（前端归档到云存储后回写） ----------------
@app.post("/api/invoices/<int:iid>/cloud-path")
@login_required("claimant", "admin")
def set_cloud_path(iid):
    u = current_user()
    body = request.get_json(silent=True) or {}
    cloud_path = (body.get("cloud_path") or "").strip() or None
    conn = db()
    r = get_invoice(conn, iid)
    if not r:
        conn.close()
        return jsonify({"ok": False, "msg": "记录不存在"}), 404
    if u["role"] == "claimant" and r["user_id"] != u["id"]:
        conn.close()
        return jsonify({"ok": False, "msg": "只能操作本人的发票"}), 403
    conn.execute("UPDATE invoices SET cloud_path=? WHERE id=?", (cloud_path, iid))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


# ---------------- 工作流动作 ----------------
@app.post("/api/invoices/<int:iid>/submit")
@login_required("claimant", "admin")
def submit_invoice(iid):
    """报销人提交：草稿/已驳回 → 待审核"""
    u = current_user()
    conn = db()
    r = get_invoice(conn, iid)
    if not r:
        conn.close()
        return jsonify({"ok": False, "msg": "记录不存在"}), 404
    if u["role"] == "claimant" and r["user_id"] != u["id"]:
        conn.close()
        return jsonify({"ok": False, "msg": "只能提交本人的发票"}), 403
    if r["status"] not in ("草稿", "已驳回"):
        conn.close()
        return jsonify({"ok": False, "msg": f"当前状态“{r['status']}”无需提交"}), 400
    conn.execute("UPDATE invoices SET status='待审核', updated_at=? WHERE id=?", (now_str(), iid))
    add_log(conn, iid, u, "提交")
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "status": "待审核"})


@app.post("/api/invoices/<int:iid>/withdraw")
@login_required("claimant", "admin")
def withdraw_invoice(iid):
    """报销人撤回：待审核 → 草稿"""
    u = current_user()
    conn = db()
    r = get_invoice(conn, iid)
    if not r:
        conn.close()
        return jsonify({"ok": False, "msg": "记录不存在"}), 404
    if u["role"] == "claimant" and r["user_id"] != u["id"]:
        conn.close()
        return jsonify({"ok": False, "msg": "只能撤回本人的发票"}), 403
    if r["status"] != "待审核":
        conn.close()
        return jsonify({"ok": False, "msg": f"当前状态“{r['status']}”不可撤回"}), 400
    conn.execute("UPDATE invoices SET status='草稿', updated_at=? WHERE id=?", (now_str(), iid))
    add_log(conn, iid, u, "撤回")
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "status": "草稿"})


@app.post("/api/invoices/<int:iid>/review")
@login_required("approver", "admin")
def review_invoice(iid):
    """审核：待审核 → 已通过 / 已驳回（附审核意见）"""
    u = current_user()
    b = request.get_json(silent=True) or {}
    action = b.get("action")
    comment = (b.get("comment") or "").strip()
    if action not in ("通过", "驳回"):
        return jsonify({"ok": False, "msg": "参数错误"}), 400
    if action == "驳回" and not comment:
        return jsonify({"ok": False, "msg": "驳回时必须填写审核意见"}), 400
    conn = db()
    r = get_invoice(conn, iid)
    if not r:
        conn.close()
        return jsonify({"ok": False, "msg": "记录不存在"}), 404
    if r["status"] != "待审核":
        conn.close()
        return jsonify({"ok": False, "msg": f"当前状态“{r['status']}”无需审核"}), 400
    new_status = "已通过" if action == "通过" else "已驳回"
    conn.execute("UPDATE invoices SET status=?, updated_at=? WHERE id=?", (new_status, now_str(), iid))
    add_log(conn, iid, u, action, comment or None)
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "status": new_status})


@app.post("/api/invoices/<int:iid>/post")
@login_required("admin")
def post_invoice(iid):
    """管理员入账：已通过 → 已入账"""
    u = current_user()
    conn = db()
    r = get_invoice(conn, iid)
    if not r:
        conn.close()
        return jsonify({"ok": False, "msg": "记录不存在"}), 404
    if r["status"] != "已通过":
        conn.close()
        return jsonify({"ok": False, "msg": "仅“已通过”的发票可入账"}), 400
    conn.execute("UPDATE invoices SET status='已入账', updated_at=? WHERE id=?", (now_str(), iid))
    add_log(conn, iid, u, "入账")
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "status": "已入账"})


@app.get("/api/invoices/<int:iid>/logs")
@login_required()
def invoice_logs(iid):
    u = current_user()
    conn = db()
    r = get_invoice(conn, iid)
    if not r:
        conn.close()
        return jsonify({"ok": False, "msg": "记录不存在"}), 404
    if u["role"] == "claimant" and r["user_id"] != u["id"]:
        conn.close()
        return jsonify({"ok": False, "msg": "没有权限查看"}), 403
    rows = conn.execute(
        "SELECT reviewer_name, action, comment, created_at FROM review_log"
        " WHERE invoice_id=? ORDER BY id DESC", (iid,)).fetchall()
    conn.close()
    return jsonify({"ok": True, "rows": [dict(x) for x in rows]})


# ---------------- 批量审核（审核人/管理员，仅对待审核记录） ----------------
@app.post("/api/invoices/batch-review")
@login_required("approver", "admin")
def batch_review():
    u = current_user()
    body = request.get_json(silent=True) or {}
    ids = body.get("ids", [])
    action = body.get("action")
    comment = (body.get("comment") or "").strip()
    if not ids or action not in ("通过", "驳回"):
        return jsonify({"ok": False, "msg": "参数错误"}), 400
    if action == "驳回" and not comment:
        return jsonify({"ok": False, "msg": "批量驳回必须填写审核意见"}), 400
    new_status = "已通过" if action == "通过" else "已驳回"
    conn = db()
    updated = skipped = 0
    for iid in ids:
        r = get_invoice(conn, iid)
        if not r or r["status"] != "待审核":
            skipped += 1
            continue
        conn.execute("UPDATE invoices SET status=?, updated_at=? WHERE id=?",
                     (new_status, now_str(), iid))
        add_log(conn, iid, u, action, comment or None)
        updated += 1
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "updated": updated, "skipped": skipped})


# ---------------- 付款登记（审核人/管理员） ----------------
PAY_METHODS = ("电汇", "承兑", "现金")


@app.post("/api/invoices/<int:iid>/payment")
@login_required("approver", "admin")
def pay_invoice(iid):
    u = current_user()
    b = request.get_json(silent=True) or {}
    paid = to_num(b.get("paid_amount"))
    pdate = (b.get("payment_date") or "").strip() or None
    method = (b.get("payment_method") or "").strip() or None
    note = (b.get("payment_note") or "").strip() or None
    if paid is None or paid < 0:
        return jsonify({"ok": False, "msg": "请填写有效的已付金额（累计）"}), 400
    if method and method not in PAY_METHODS:
        return jsonify({"ok": False, "msg": f"付款方式仅支持：{'/'.join(PAY_METHODS)}"}), 400
    conn = db()
    r = get_invoice(conn, iid)
    if not r:
        conn.close()
        return jsonify({"ok": False, "msg": "记录不存在"}), 404
    conn.execute(
        "UPDATE invoices SET paid_amount=?, payment_date=?, payment_method=?, payment_note=?, updated_at=? WHERE id=?",
        (paid, pdate, method, note, now_str(), iid))
    status, bal = payment_info(r["total_amount"], paid)
    bal_show = f"{bal:,.2f}" if bal is not None else "-"
    log_txt = f"付款日期 {pdate or '-'}，方式 {method or '-'}，累计已付 ¥{paid:,.2f}，未付差额 ¥{bal_show}（{status}）"
    if note:
        log_txt += f"；备注：{note}"
    add_log(conn, iid, u, "付款", log_txt)
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "paid_amount": paid, "payment_date": pdate,
                    "payment_method": method, "payment_note": note,
                    "payment_status": status, "balance": bal})


# ---------------- 按单位付款/收款（审核人/管理员） ----------------
# mode: seller=付销售方(进项) / claimant=付报销人(进项) / buyer=收客户款(销项)
def _unpaid_invoices(conn, col, name, direction, payee=None):
    """查询单位名下未付清/未收清的发票（仅做账、已驳回除外），按开票顺序返回"""
    sql = f"""SELECT id, invoice_no, invoice_date, total_amount, COALESCE(paid_amount,0) paid
            FROM invoices WHERE {col}=? AND direction=? AND settle='正常付款'
              AND status!='已驳回' AND total_amount IS NOT NULL"""
    params = [name, direction]
    if payee:
        sql += " AND payee=?"
        params.append(payee)
    sql += " ORDER BY invoice_date, id LIMIT 200"
    rows = conn.execute(sql, params).fetchall()
    out = []
    for r in rows:
        st, bal = payment_info(r["total_amount"], r["paid"])
        if bal is not None and bal > 0.004:
            d = dict(r)
            d["balance"] = bal
            out.append(d)
    return out


def _pay_mode_conf(mode):
    """mode → (列名, direction, 是否带payee, 标签)"""
    return {"seller": ("seller_name", "进项", "销售方", "销售方"),
            "claimant": ("claimant", "进项", "报销人", "报销人"),
            "buyer": ("buyer_name", "销项", None, "客户")}[mode]


@app.get("/api/pay-unpaid")
@login_required("approver", "admin")
def pay_unpaid():
    mode = (request.args.get("mode") or "").strip()
    name = (request.args.get("name") or "").strip()
    if mode not in ("seller", "claimant", "buyer") or not name:
        return jsonify({"ok": False, "msg": "参数错误"}), 400
    col, direction, payee, _ = _pay_mode_conf(mode)
    conn = db()
    rows = _unpaid_invoices(conn, col, name, direction, payee)
    conn.close()
    return jsonify({"ok": True, "rows": rows})


@app.post("/api/pay-batch")
@login_required("approver", "admin")
def pay_batch():
    """整笔付款/收款按开票顺序自动抵扣某单位名下未付清的发票"""
    u = current_user()
    b = request.get_json(silent=True) or {}
    mode = (b.get("mode") or "").strip()
    name = (b.get("name") or "").strip()
    amount = to_num(b.get("amount"))
    pdate = (b.get("payment_date") or "").strip() or None
    method = (b.get("payment_method") or "").strip() or None
    note = (b.get("payment_note") or "").strip() or None
    if mode not in ("seller", "claimant", "buyer") or not name:
        return jsonify({"ok": False, "msg": "参数错误"}), 400
    if amount is None or amount <= 0:
        return jsonify({"ok": False, "msg": "请填写有效的付款金额"}), 400
    if method and method not in PAY_METHODS:
        return jsonify({"ok": False, "msg": f"付款方式仅支持：{'/'.join(PAY_METHODS)}"}), 400
    col, direction, payee, tag = _pay_mode_conf(mode)
    label = f"{tag} {name}"
    verb = "收款" if mode == "buyer" else "付款"
    conn = db()
    rows = _unpaid_invoices(conn, col, name, direction, payee)
    if not rows:
        conn.close()
        return jsonify({"ok": False, "msg": f"{label} 名下没有未付清的发票"}), 400
    remain = round(amount, 2)
    applied = []
    for r in rows:
        if remain <= 0.004:
            break
        pay = min(r["balance"], remain)
        new_paid = round(r["paid"] + pay, 2)
        status, _ = payment_info(r["total_amount"], new_paid)
        log_note = (note + "；" if note else "") + f"[{label} 整笔{verb} ¥{amount:,.2f} 按序抵扣 ¥{pay:,.2f}]"
        conn.execute(
            "UPDATE invoices SET paid_amount=?, payment_date=?, payment_method=?, payment_note=?, updated_at=? WHERE id=?",
            (new_paid, pdate, method, log_note, now_str(), r["id"]))
        add_log(conn, r["id"], u, "付款",
                f"{label} 整笔{verb} ¥{amount:,.2f}（{method or '-'}，{pdate or '-'}），本票抵扣 ¥{pay:,.2f}，"
                f"累计{'已收' if mode == 'buyer' else '已付'} ¥{new_paid:,.2f}（{status}）" + (f"；备注：{note}" if note else ""))
        applied.append({"id": r["id"], "invoice_no": r["invoice_no"], "pay": round(pay, 2)})
        remain = round(remain - pay, 2)
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "applied": applied, "count": len(applied),
                    "unallocated": remain if remain > 0.004 else 0,
                    "msg": f"{label}：已抵扣 {len(applied)} 张发票" + (f"，剩余 ¥{remain:,.2f} 无票可抵" if remain > 0.004 else "")})


# ---------------- 付款记录（查询/导出，审核人/管理员） ----------------
def _payment_where(q):
    """付款记录筛选条件：有付款记录（paid_amount>0）的发票"""
    where = ["paid_amount IS NOT NULL", "paid_amount > 0"]
    params = []
    kw = (q.get("keyword") or "").strip()
    if kw:
        like = f"%{kw}%"
        where.append("(invoice_no LIKE ? OR seller_name LIKE ? OR claimant LIKE ? OR payment_note LIKE ?)")
        params += [like] * 4
    date_from = (q.get("date_from") or "").strip()
    date_to = (q.get("date_to") or "").strip()
    if date_from:
        where.append("payment_date >= ?")
        params.append(date_from)
    if date_to:
        where.append("payment_date <= ?")
        params.append(date_to)
    method = (q.get("method") or "").strip()
    if method:
        where.append("payment_method = ?")
        params.append(method)
    payee = (q.get("payee") or "").strip()
    if payee in PAYEES:
        where.append("payee = ?")
        params.append(payee)
    status = (q.get("status") or "").strip()
    if status in ("已付清", "部分付款", "未付款"):
        if status == "未付款":
            where.append("0=1")  # 有付款记录的不存在未付款
        else:
            where.append("CASE WHEN total_amount - paid_amount > 0.005 THEN '部分付款' ELSE '已付清' END = ?")
            params.append(status)
    return " AND ".join(where), params


_PAY_COLS = ("id, invoice_no, invoice_type, invoice_date, seller_name, claimant, payee, settle,"
             " total_amount, paid_amount, payment_date, payment_method, payment_note, status")


@app.get("/api/payments")
@login_required("approver", "admin")
def payments_list():
    cond, params = _payment_where(request.args)
    conn = db()
    rows = conn.execute(
        f"SELECT {_PAY_COLS} FROM invoices WHERE {cond} ORDER BY payment_date DESC, id DESC LIMIT 1000",
        params).fetchall()
    agg = conn.execute(
        f"SELECT COUNT(*) n, SUM(total_amount) t, SUM(paid_amount) p FROM invoices WHERE {cond}",
        params).fetchone()
    conn.close()
    out = []
    for r in rows:
        d = dict(r)
        d["payment_status"], d["balance"] = payment_info(d["total_amount"], d["paid_amount"])
        out.append(d)
    return jsonify({"ok": True, "total": agg["n"] or 0,
                    "sum_total": round(agg["t"] or 0, 2), "sum_paid": round(agg["p"] or 0, 2),
                    "rows": out})


@app.get("/api/payments-export")
@login_required("approver", "admin")
def payments_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    cond, params = _payment_where(request.args)
    conn = db()
    rows = conn.execute(
        f"SELECT {_PAY_COLS} FROM invoices WHERE {cond} ORDER BY payment_date DESC, id DESC", params).fetchall()
    conn.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "付款记录"
    headers = ["付款日期", "付款方式", "收款对象", "收款单位/人", "发票号码", "发票类型",
               "开票日期", "价税合计", "已付金额", "未付差额", "付款状态", "备注", "台账编号", "发票状态"]
    ws.append(headers)
    hf = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2563EB")
    for c in ws[1]:
        c.font, c.fill, c.alignment = hf, fill, Alignment(horizontal="center")
    n = 0
    for r in rows:
        d = dict(r)
        st, bal = payment_info(d["total_amount"], d["paid_amount"])
        ws.append([d["payment_date"], d["payment_method"] or "",
                   d["payee"], d["claimant"] if d["payee"] == "报销人" else d["seller_name"],
                   d["invoice_no"], d["invoice_type"], d["invoice_date"],
                   d["total_amount"], d["paid_amount"], bal, st,
                   d["payment_note"], d["id"], d["status"]])
        n += 1
    if n:
        last = ws.max_row
        ws.append(["合计", "", "", "", "", "", "",
                   f"=SUM(H2:H{last})", f"=SUM(I2:I{last})", f"=SUM(J2:J{last})", "", "", "", ""])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    widths = [12, 10, 10, 26, 24, 18, 12, 14, 14, 14, 10, 26, 8, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r_ in range(2, ws.max_row + 1):
        for col in (8, 9, 10):
            ws.cell(row=r_, column=col).number_format = "#,##0.00"
    tag = f"{request.args.get('date_from') or ''}~{request.args.get('date_to') or ''}" if (request.args.get('date_from') or request.args.get('date_to')) else datetime.now().strftime("%Y%m%d")
    out = os.path.join(BASE_DIR, "付款记录导出.xlsx")
    wb.save(out)
    return send_file(out, as_attachment=True, download_name=f"付款记录_{tag}.xlsx")


# ---------------- 进销项增值税统计（审核人/管理员） ----------------
# 进项税额只统计增值税专用发票（可抵扣）；普票/其他发票税额不计入
_DEDUCTIBLE = "invoice_type LIKE '%专用%'"


@app.get("/api/vat-stats")
@login_required("approver", "admin")
def vat_stats():
    year = (request.args.get("year") or "").strip()
    month = (request.args.get("month") or "").strip()
    company_id = (request.args.get("company_id") or "").strip()
    date_from = (request.args.get("date_from") or "").strip()
    date_to = (request.args.get("date_to") or "").strip()
    range_mode = bool(date_from or date_to)
    if not range_mode:
        if not re.match(r"^\d{4}$", year):
            year = str(datetime.now().year)
        month = month.zfill(2) if re.match(r"^\d{1,2}$", month) else ""

    where, params = ["direction IS NOT NULL"], []
    if range_mode:
        if date_from:
            where.append("invoice_date >= ?")
            params.append(date_from)
        if date_to:
            where.append("invoice_date <= ?")
            params.append(date_to)
    else:
        where.append("invoice_date LIKE ?")
        params.append(year + "-%")
        if month:
            where.append("substr(invoice_date,6,2)=?")
            params.append(month)
    if company_id.isdigit():
        where.append("company_id=?")
        params.append(int(company_id))
    cond = " AND ".join(where)

    agg_sql = (
        "SUM(CASE WHEN direction='进项' THEN 1 ELSE 0 END) in_n,"
        "SUM(CASE WHEN direction='进项' THEN total_amount ELSE 0 END) in_total,"
        f"SUM(CASE WHEN direction='进项' AND {_DEDUCTIBLE} THEN tax_amount ELSE 0 END) in_tax,"
        "SUM(CASE WHEN direction='进项' THEN amount ELSE 0 END) in_amount,"
        f"SUM(CASE WHEN direction='进项' AND NOT COALESCE({_DEDUCTIBLE},0) THEN tax_amount ELSE 0 END) in_tax_other,"
        "SUM(CASE WHEN direction='销项' THEN 1 ELSE 0 END) out_n,"
        "SUM(CASE WHEN direction='销项' THEN total_amount ELSE 0 END) out_total,"
        "SUM(CASE WHEN direction='销项' THEN tax_amount ELSE 0 END) out_tax,"
        "SUM(CASE WHEN direction='销项' THEN amount ELSE 0 END) out_amount")

    conn = db()
    s = conn.execute(f"SELECT {agg_sql} FROM invoices WHERE {cond}", params).fetchone()
    months = conn.execute(
        f"SELECT substr(invoice_date,1,7) m, {agg_sql} FROM invoices"
        f" WHERE {cond} GROUP BY m ORDER BY m", params).fetchall()

    # 年份下拉：从发票开票日期自动取年份（含当年，供未来年份预置）
    years = [r[0] for r in conn.execute(
        "SELECT DISTINCT substr(invoice_date,1,4) y FROM invoices"
        " WHERE invoice_date IS NOT NULL AND length(invoice_date)>=4 ORDER BY y DESC")]
    cur_year = str(datetime.now().year)
    if cur_year not in years:
        years.append(cur_year)

    def details(direction):
        rr = conn.execute(
            f"""SELECT id, invoice_no, invoice_type, invoice_date, buyer_name, seller_name, item_name,
                       amount, tax_rate, tax_amount, total_amount, status, payment_date,
                       paid_amount, payment_method
                FROM invoices WHERE {cond} AND direction=?
                ORDER BY invoice_date DESC, id DESC LIMIT 500""",
            params + [direction]).fetchall()
        return [dict(x) for x in rr]

    in_rows, out_rows = details("进项"), details("销项")
    conn.close()

    def f(v):
        return round(v or 0, 2)

    return jsonify({
        "ok": True, "year": year, "month": month,
        "date_from": date_from, "date_to": date_to,
        "years": years,
        "in_n": s["in_n"] or 0, "in_amount": f(s["in_amount"]),
        "in_tax": f(s["in_tax"]), "in_tax_other": f(s["in_tax_other"]),
        "in_total": f(s["in_total"]),
        "out_n": s["out_n"] or 0, "out_amount": f(s["out_amount"]),
        "out_tax": f(s["out_tax"]), "out_total": f(s["out_total"]),
        "diff_tax": round((s["out_tax"] or 0) - (s["in_tax"] or 0), 2),
        "diff_total": round((s["out_total"] or 0) - (s["in_total"] or 0), 2),
        "months": [dict(x) for x in months],
        "in_rows": in_rows, "out_rows": out_rows,
    })


# ---------------- 进销项发票明细导出（按日期区间） ----------------
@app.get("/api/vat-export")
@login_required("approver", "admin")
def vat_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    direction = request.args.get("direction") or "进项"
    if direction not in ("进项", "销项"):
        return jsonify({"ok": False, "msg": "direction 参数须为 进项/销项"}), 400
    date_from = (request.args.get("date_from") or "").strip()
    date_to = (request.args.get("date_to") or "").strip()
    company_id = (request.args.get("company_id") or "").strip()

    where, params = ["direction=?", "invoice_date IS NOT NULL"], [direction]
    if date_from:
        where.append("invoice_date >= ?")
        params.append(date_from)
    if date_to:
        where.append("invoice_date <= ?")
        params.append(date_to)
    if company_id.isdigit():
        where.append("company_id=?")
        params.append(int(company_id))
    cond = " AND ".join(where)

    conn = db()
    rows = conn.execute(
        f"""SELECT invoice_no, invoice_type, invoice_date, buyer_name, seller_name, item_name,
                   amount, tax_rate, tax_amount, total_amount, status,
                   payment_date, payment_method, paid_amount
            FROM invoices WHERE {cond} ORDER BY invoice_date, invoice_no""", params).fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = f"{direction}发票明细"
    headers = ["开票日期", "发票号码", "发票类型", "销售方", "购买方", "品名",
               "不含税金额", "税率", "税额", "价税合计", "状态",
               "付款日期", "付款方式", "已付金额"]
    ws.append(headers)
    hf = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2563EB")
    for c in ws[1]:
        c.font, c.fill, c.alignment = hf, fill, Alignment(horizontal="center")
    for r in rows:
        d = dict(r)
        ws.append([d["invoice_date"], d["invoice_no"], d["invoice_type"] or "",
                   d["seller_name"], d["buyer_name"], d["item_name"],
                   d["amount"], d["tax_rate"], d["tax_amount"], d["total_amount"],
                   d["status"], d.get("payment_date"), d.get("payment_method"),
                   d.get("paid_amount")])
    widths = [12, 24, 18, 26, 26, 26, 14, 8, 12, 14, 10, 12, 10, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    last = ws.max_row
    # 合计行
    ws.append(["合计", "", "", "", "", "",
               f"=SUM(G2:G{last})" if last > 1 else "", "", f"=SUM(I2:I{last})" if last > 1 else "",
               f"=SUM(J2:J{last})" if last > 1 else "", "", "", "", f"=SUM(N2:N{last})" if last > 1 else ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r_ in range(2, ws.max_row + 1):
        for col in (7, 9, 10, 14):
            ws.cell(row=r_, column=col).number_format = "#,##0.00"
    tag = f"{date_from or ''}~{date_to or ''}" if (date_from or date_to) else datetime.now().strftime("%Y%m%d")
    out = os.path.join(BASE_DIR, f"{direction}发票明细导出.xlsx")
    wb.save(out)
    return send_file(out, as_attachment=True, download_name=f"{direction}发票明细_{tag}.xlsx")


# ---------------- 批量状态（仅管理员） ----------------
@app.post("/api/invoices/batch-status")
@login_required("admin")
def batch_status():
    u = current_user()
    body = request.get_json(silent=True) or {}
    ids, status = body.get("ids", []), body.get("status")
    if not ids or status not in {"待审核", "已通过", "已驳回", "已入账", "草稿"}:
        return jsonify({"ok": False, "msg": "参数错误"}), 400
    conn = db()
    qs = ",".join("?" for _ in ids)
    cur = conn.execute(f"UPDATE invoices SET status=?, updated_at=? WHERE id IN ({qs})",
                       [status, now_str()] + list(ids))
    for iid in ids:
        add_log(conn, iid, u, "批量", f"→ {status}")
    conn.commit()
    n = cur.rowcount
    conn.close()
    return jsonify({"ok": True, "updated": n})


# ---------------- 统计（按角色范围） ----------------
@app.get("/api/stats")
@login_required()
def stats():
    u = current_user()
    scope, params = "1=1", []
    if u["role"] == "claimant":
        scope, params = "user_id = ?", [u["id"]]
    elif u["role"] == "approver":
        scope = "status != '草稿'"
    # 按公司抬头筛选（收票付款 / 开票收款 / 总览卡片同步生效）
    company_id = (request.args.get("company_id") or "").strip()
    if company_id.isdigit():
        scope += " AND company_id=?"
        params.append(int(company_id))
    conn = db()
    overall = conn.execute(
        f"SELECT COUNT(*) n, SUM(total_amount) t, SUM(tax_amount) x FROM invoices WHERE {scope}", params).fetchone()
    by_subject = conn.execute(
        f"SELECT subject_code, subject_name, COUNT(*) n, SUM(total_amount) t "
        f"FROM invoices WHERE {scope} GROUP BY subject_code, subject_name ORDER BY t DESC", params).fetchall()
    by_month = conn.execute(
        f"SELECT substr(invoice_date,1,7) m, COUNT(*) n, SUM(total_amount) t "
        f"FROM invoices WHERE invoice_date IS NOT NULL AND {scope} GROUP BY m ORDER BY m DESC LIMIT 12", params).fetchall()
    by_status = conn.execute(
        f"SELECT status, COUNT(*) n, SUM(total_amount) t FROM invoices WHERE {scope} GROUP BY status", params).fetchall()
    by_type = conn.execute(
        f"SELECT invoice_type, COUNT(*) n, SUM(tax_amount) x FROM invoices WHERE {scope} GROUP BY invoice_type", params).fetchall()
    # 收票付款统计：仅做账、已驳回的发票不计入欠款；按收款对象（销售方/报销人）分组
    pay_base = (f"{scope} AND direction='进项' AND total_amount IS NOT NULL"
                f" AND settle='正常付款' AND status!='已驳回'")
    pay_by_seller = conn.execute(
        f"""SELECT seller_name, COUNT(*) n, SUM(total_amount) t,
                   SUM(COALESCE(paid_amount,0)) p, MAX(payment_date) last_pay
            FROM invoices WHERE {pay_base} AND payee='销售方' GROUP BY seller_name ORDER BY t DESC""", params).fetchall()
    pay_by_claimant = conn.execute(
        f"""SELECT claimant, COUNT(*) n, SUM(total_amount) t,
                   SUM(COALESCE(paid_amount,0)) p, MAX(payment_date) last_pay
            FROM invoices WHERE {pay_base} AND payee='报销人' GROUP BY claimant ORDER BY t DESC""", params).fetchall()
    pay_all = conn.execute(
        f"SELECT COUNT(*) n, SUM(total_amount) t, SUM(COALESCE(paid_amount,0)) p FROM invoices WHERE {pay_base}",
        params).fetchone()
    book_all = conn.execute(
        f"SELECT COUNT(*) n, COALESCE(SUM(total_amount),0) t FROM invoices"
        f" WHERE {scope} AND direction='进项' AND total_amount IS NOT NULL AND settle='仅做账'", params).fetchone()
    rej_all = conn.execute(
        f"SELECT COUNT(*) n, COALESCE(SUM(total_amount),0) t FROM invoices"
        f" WHERE {scope} AND direction='进项' AND total_amount IS NOT NULL AND settle='正常付款' AND status='已驳回'",
        params).fetchone()
    # 开票收款统计：我方开出的销项发票，按客户统计收款情况
    rcpt_base = (f"{scope} AND direction='销项' AND total_amount IS NOT NULL"
                 f" AND settle='正常付款' AND status!='已驳回'")
    receipt_by_buyer = conn.execute(
        f"""SELECT buyer_name, COUNT(*) n, SUM(total_amount) t,
                   SUM(COALESCE(paid_amount,0)) p, MAX(payment_date) last_pay
            FROM invoices WHERE {rcpt_base} GROUP BY buyer_name ORDER BY t DESC""", params).fetchall()
    receipt_all = conn.execute(
        f"SELECT COUNT(*) n, SUM(total_amount) t, SUM(COALESCE(paid_amount,0)) p FROM invoices WHERE {rcpt_base}",
        params).fetchone()
    # 总览卡片：有效发票口径（正常付款、非驳回）
    eff_base = f"{scope} AND settle='正常付款' AND status!='已驳回' AND total_amount IS NOT NULL"
    eff_in = conn.execute(
        f"""SELECT COUNT(*) n, SUM(total_amount) t,
                   SUM(CASE WHEN invoice_type LIKE '%专用%' THEN tax_amount ELSE 0 END) tax_deduct,
                   SUM(COALESCE(tax_amount,0)) tax_all
            FROM invoices WHERE {eff_base} AND direction='进项'""", params).fetchone()
    eff_out = conn.execute(
        f"""SELECT COUNT(*) n, SUM(total_amount) t, SUM(COALESCE(tax_amount,0)) tax_all
            FROM invoices WHERE {eff_base} AND direction='销项'""", params).fetchone()
    conn.close()
    return jsonify({
        "ok": True,
        "count": overall["n"], "total": overall["t"] or 0, "tax": overall["x"] or 0,
        "by_subject": [dict(r) for r in by_subject],
        "by_month": [dict(r) for r in by_month],
        "by_status": [dict(r) for r in by_status],
        "by_type": [dict(r) for r in by_type],
        "pay_by_seller": [dict(r) for r in pay_by_seller],
        "pay_by_claimant": [dict(r) for r in pay_by_claimant],
        "pay_all": {"n": pay_all["n"] or 0, "total": round(pay_all["t"] or 0, 2),
                    "paid": round(pay_all["p"] or 0, 2),
                    "owed": round((pay_all["t"] or 0) - (pay_all["p"] or 0), 2)},
        "book_all": {"n": book_all["n"] or 0, "total": round(book_all["t"] or 0, 2)},
        "rej_all": {"n": rej_all["n"] or 0, "total": round(rej_all["t"] or 0, 2)},
        "receipt_by_buyer": [dict(r) for r in receipt_by_buyer],
        "receipt_all": {"n": receipt_all["n"] or 0, "total": round(receipt_all["t"] or 0, 2),
                        "received": round(receipt_all["p"] or 0, 2),
                        "owed": round((receipt_all["t"] or 0) - (receipt_all["p"] or 0), 2)},
        "eff_in": {"n": eff_in["n"] or 0, "total": round(eff_in["t"] or 0, 2),
                   "tax_deduct": round(eff_in["tax_deduct"] or 0, 2), "tax_all": round(eff_in["tax_all"] or 0, 2)},
        "eff_out": {"n": eff_out["n"] or 0, "total": round(eff_out["t"] or 0, 2),
                    "tax_all": round(eff_out["tax_all"] or 0, 2)},
    })


# ---------------- 导出 Excel（审核人/管理员） ----------------
@app.get("/api/export")
@login_required("approver", "admin")
def export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    conn = db()
    rows = conn.execute("SELECT * FROM invoices ORDER BY invoice_date DESC, id DESC").fetchall()
    conn.close()
    headers = ["台账编号", "发票号码", "发票代码", "发票类型", "开票日期", "购买方", "销售方",
               "品名", "不含税金额", "税率", "税额", "价税合计", "会计科目", "报销人", "部门",
               "项目", "状态", "结算", "付款日期", "付款方式", "已付金额", "未付差额", "备注", "录入时间"]
    wb = Workbook()
    ws = wb.active
    ws.title = "发票台账"
    ws.append(headers)
    hf = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2563EB")
    for c in ws[1]:
        c.font, c.fill, c.alignment = hf, fill, Alignment(horizontal="center")
    for r in rows:
        d = dict(r)
        _st, _bal = payment_info(d["total_amount"], d.get("paid_amount"))
        settle = "仅做账" if d.get("settle") == "仅做账" else ("付报销人" if d.get("payee") == "报销人" else "")
        ws.append([d["id"], d["invoice_no"], d["invoice_code"], d["invoice_type"], d["invoice_date"],
                   d["buyer_name"], d["seller_name"], d["item_name"], d["amount"], d["tax_rate"],
                   d["tax_amount"], d["total_amount"],
                   f"{d['subject_code'] or ''} {d['subject_name'] or ''}".strip(),
                   d["claimant"], d["department"], d["project"], d["status"], settle,
                   d.get("payment_date"), d.get("payment_method"), d.get("paid_amount"), _bal,
                   d["remark"], d["created_at"]])
    widths = [8, 22, 12, 18, 12, 22, 22, 24, 12, 8, 10, 12, 20, 10, 10, 14, 10, 10, 12, 10, 12, 12, 18, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r_ in range(2, ws.max_row + 1):
        for col in (9, 11, 12, 21, 22):
            ws.cell(row=r_, column=col).number_format = "#,##0.00"
    out = os.path.join(BASE_DIR, "发票台账导出.xlsx")
    wb.save(out)
    return send_file(out, as_attachment=True, download_name=f"发票台账_{datetime.now():%Y%m%d}.xlsx")


# ---------------- 发票原件下载 ----------------
@app.get("/api/files/<path:fname>")
@login_required()
def get_file(fname):
    u = current_user()
    base = os.path.basename(fname)
    conn = db()
    r = conn.execute("SELECT user_id FROM invoices WHERE file_path=?", (base,)).fetchone()
    conn.close()
    if r and u["role"] == "claimant" and r["user_id"] != u["id"]:
        return jsonify({"ok": False, "msg": "没有权限下载该原件"}), 403
    return send_from_directory(UPLOAD_DIR, base, as_attachment=True)


# ---------------- 数据库备份 / 恢复（仅管理员） ----------------
REQUIRED_TABLES = {"invoices", "users", "companies", "review_log"}


def snapshot_db(dest_path):
    """在线一致性快照（sqlite3 backup API，自动合并 WAL，不阻塞服务）"""
    src = sqlite3.connect(DB_PATH)
    try:
        dst = sqlite3.connect(dest_path)
        try:
            src.backup(dst)
        finally:
            dst.close()
    finally:
        src.close()


@app.get("/api/backup/download")
@login_required("admin")
def backup_download():
    snap = os.path.join(BASE_DIR, f"~backup_{int(time.time() * 1000)}.db")
    snapshot_db(snap)
    dl_name = f"发票与付款系统备份_{datetime.now():%Y%m%d_%H%M%S}.db"

    @after_this_request
    def _cleanup(resp):
        try:
            os.remove(snap)
        except OSError:
            pass
        return resp

    return send_file(snap, as_attachment=True, download_name=dl_name)


@app.post("/api/backup/restore")
@login_required("admin")
def backup_restore():
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify({"ok": False, "msg": "未收到备份文件"}), 400
    tmp = os.path.join(BASE_DIR, f"~restore_{int(time.time() * 1000)}.db")
    f.save(tmp)

    # ---- 校验备份文件：完整性 + 必需数据表 ----
    try:
        chk = sqlite3.connect(tmp)
        try:
            row = chk.execute("PRAGMA integrity_check").fetchone()
            if not row or row[0] != "ok":
                raise ValueError("数据库完整性校验失败")
            tables = {r[0] for r in
                      chk.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            missing = REQUIRED_TABLES - tables
            if missing:
                raise ValueError("缺少数据表：" + "、".join(sorted(missing)))
        finally:
            chk.close()
    except Exception as e:
        _trash(tmp)
        return jsonify({"ok": False, "msg": f"备份文件无效：{e}"}), 422

    # ---- 替换当前库：先 checkpoint WAL，再把现库留一份 .bak ----
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = os.path.join(BASE_DIR, f"reimbursement.db.before-restore-{ts}.bak")
    try:
        c = sqlite3.connect(DB_PATH)
        try:
            c.execute("PRAGMA wal_checkpoint(TRUNCATE)")
        finally:
            c.close()
        shutil.copy2(DB_PATH, bak)
        for ext in ("-wal", "-shm"):
            p = DB_PATH + ext
            if os.path.exists(p):
                os.remove(p)
        os.replace(tmp, DB_PATH)
    except OSError as e:
        return jsonify({"ok": False, "msg": f"恢复失败：{e}（原数据库已保留）"}), 500

    # ---- 在新库上跑迁移/索引，并重算进销项方向 ----
    init_db()
    conn = db()
    recompute_all_directions(conn)
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "msg": "恢复成功", "backup_file": os.path.basename(bak)})


if __name__ == "__main__":
    print("发票与付款系统已启动: http://127.0.0.1:8686")
    app.run(host="127.0.0.1", port=8686, debug=False)
